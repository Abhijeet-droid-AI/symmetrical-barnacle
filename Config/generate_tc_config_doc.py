"""
Generates 'Classification Utilities - TC Config Integration Guide.docx' in the
project root. Documents the CORRECTED Teamcenter configuration split:

    - classification_utilities.cfg  : ALL library + function settings,
                                      including [ENVIRONMENT] TC_ROOT / TC_DATA
                                      (developer-owned - users never edit it)
    - tc_config.txt (NEW)           : Teamcenter LOGIN ONLY
                                      (TC_USER / TC_PASS / TC_GROUP)
                                      (user-owned - the ONLY file a user edits)

This REWORKS the earlier attempt that had moved TC_ROOT / TC_DATA together
with the credentials into tc_environment.cfg. That was wrong: the Teamcenter
library settings belong to the developer side. Now Build_All.bat reads TC_ROOT
from the master cfg again, the Run_*.bat wrappers read TC_ROOT / TC_DATA from
the master cfg and only the login from tc_config.txt, and every exe resolves
its credentials from tc_config.txt (directly or via -tcconfig=).

Structure: title page, why the previous split was reworked, architecture,
change-by-change sections with copy-paste ready code, delivery + user
workflow, build procedure, verification checklist, troubleshooting, rollback.

Run:  python Config/generate_tc_config_doc.py
"""

import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_PATH = os.path.join(
    PROJECT_ROOT, "Classification Utilities - TC Config Integration Guide.docx")

MONO = "Courier New"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
GREEN = RGBColor(0x00, 0x66, 0x00)
AMBER = RGBColor(0x8A, 0x53, 0x00)
RED = RGBColor(0xA0, 0x10, 0x10)

doc = Document()


# ----------------------------------------------------------------- helpers
def shade(cell, hex_fill):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(el)


def add_code(text, font_size=8.5):
    """Code block rendered as a one-cell table with a light background."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    shade(cell, "F2F2F2")
    cell.paragraphs[0].text = ""
    lines = text.strip("\n").split("\n")
    for idx, line in enumerate(lines):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line)
        run.font.name = MONO
        run.font.size = Pt(font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def h1(text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        r.font.color.rgb = ACCENT


def note(text):
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.rows[0].cells[0]
    shade(cell, "EAF1FB")
    cell.paragraphs[0].text = ""
    run = cell.paragraphs[0].add_run(text)
    run.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def warn(text):
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.rows[0].cells[0]
    shade(cell, "FFF2CC")
    cell.paragraphs[0].text = ""
    run = cell.paragraphs[0].add_run(text)
    run.font.size = Pt(9.5)
    run.font.color.rgb = AMBER
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def bad(text):
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.rows[0].cells[0]
    shade(cell, "FBE4E4")
    cell.paragraphs[0].text = ""
    run = cell.paragraphs[0].add_run(text)
    run.font.size = Pt(9.5)
    run.font.color.rgb = RED
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def expected(text):
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.rows[0].cells[0]
    shade(cell, "E9F0E4")
    cell.paragraphs[0].text = ""
    run = cell.paragraphs[0].add_run("Expected result: " + text)
    run.font.size = Pt(9.5)
    run.font.color.rgb = GREEN
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def bullets(lines):
    for line in lines:
        doc.add_paragraph(line, style="List Bullet")


def numbered(lines):
    for i, line in enumerate(lines, start=1):
        doc.add_paragraph("%d. %s" % (i, line))


def grid_table(hdr, rows, widths):
    tbl = doc.add_table(rows=len(rows) + 1, cols=len(hdr))
    tbl.style = "Table Grid"
    for i, htxt in enumerate(hdr):
        c = tbl.rows[0].cells[i]
        c.paragraphs[0].text = ""
        run = c.paragraphs[0].add_run(htxt)
        run.bold = True
        shade(c, "D9E2F3")
    for ri, vals in enumerate(rows, start=1):
        for ci, val in enumerate(vals):
            c = tbl.rows[ri].cells[ci]
            c.paragraphs[0].text = ""
            run = c.paragraphs[0].add_run(val)
            run.font.size = Pt(8.5)
            if ci == 0:
                run.bold = True
                shade(c, "EAF1FB")
    for w, col in zip(widths, range(len(hdr))):
        for cell in tbl.columns[col].cells:
            cell.width = w


# ------------------------------------------------------------- title page
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Classification Utilities TC13")
r.bold = True
r.font.size = Pt(26)

t2 = doc.add_paragraph()
t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t2.add_run("Teamcenter Credential Split (tc_config.txt) - Integration Guide")
r.bold = True
r.font.size = Pt(20)
r.font.color.rgb = ACCENT

t3 = doc.add_paragraph()
t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t3.add_run(
    "ALL Teamcenter library calls and settings (TC_ROOT / TC_DATA) stay inside\n"
    "classification_utilities.cfg. Only the TC LOGIN moves to a separate\n"
    "user-owned file tc_config.txt. After a Build_All.bat the exes in\n"
    "x64\\Release are delivered; the user edits ONLY tc_config.txt, then runs\n"
    "the exes via Run.bat / Run_*.bat - no other file, no cfg edits, ever.")
r.font.size = Pt(11)

t4 = doc.add_paragraph()
t4.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t4.add_run("Scope: ClassificationExtraction / ClassificationImport / "
               "ClassificationDelete - Config bat files - master cfg file")
r.font.size = Pt(9.5)

doc.add_paragraph()
note("HOW TO USE THIS DOCUMENT: section 2 explains why the earlier split was "
     "reworked, section 3 is the new architecture, sections 5-9 describe every "
     "file change with code, section 10 is the delivery package + the user "
     "workflow, section 11 is the developer build procedure, section 12 is the "
     "verification checklist, section 13 is troubleshooting, section 14 is "
     "rollback.")

# ------------------------------------------------------------- 1. summary
h1("1. Requirement and Solution Summary")
p = doc.add_paragraph(
    "Requirement: all Teamcenter library configuration (TC_ROOT, TC_DATA, "
    "include/lib paths used by every ITK call) must stay in the developer "
    "config classification_utilities.cfg. ONLY the Teamcenter login "
    "credentials must live in a separate file, tc_config.txt, which the user "
    "edits once after receiving the exes. When the user runs a utility, the "
    "exe must pick up those credentials and perform its task - with zero "
    "changes to any other config file on the user's machine.")
grid_table(
    ["", "Before (wrong attempt)", "After (this rework)"],
    [
        ("TC_ROOT / TC_DATA (library)",
         "[ENVIRONMENT] in tc_environment.cfg - user had to edit it",
         "[ENVIRONMENT] in classification_utilities.cfg - developer-owned"),
        ("Teamcenter login",
         "[CREDENTIALS] in tc_environment.cfg",
         "[CREDENTIALS] in tc_config.txt (NEW file, user-owned)"),
        ("tc_environment.cfg",
         "held everything TC-specific",
         "DELETED - replaced by tc_config.txt"),
        ("What the user edits",
         "tc_environment.cfg (paths AND login)",
         "tc_config.txt ONLY (login only)"),
        ("What the developer edits",
         "classification_utilities.cfg",
         "classification_utilities.cfg (unchanged responsibility)"),
    ],
    (Inches(1.6), Inches(2.5), Inches(2.5)))

doc.add_paragraph()
doc.add_paragraph(
    "The function configuration ([LOGS], [EXTRACTION], [IMPORT], [DELETE], "
    "[RUN], [BUILD]) stays in classification_utilities.cfg exactly as before. "
    "The exes keep the -tcconfig= argument; it now points at tc_config.txt.")

# ------------------------------------------------------------- 2. why rework
h1("2. Why the Previous Split Was Reworked")
bad("The previous change moved TC_ROOT and TC_DATA together with the "
    "credentials into tc_environment.cfg. That made the USER responsible for "
    "Teamcenter library settings: a wrong TC_ROOT broke both the build and "
    "the tc_profilevars.bat runtime setup, and the delivered package was not "
    "usable without editing developer-owned values. TC_ROOT / TC_DATA belong "
    "to the developer's build and runtime environment - only the login "
    "belongs to the user.")
doc.add_paragraph("Correction applied in this rework:")
bullets([
    "[ENVIRONMENT] TC_ROOT / TC_DATA restored in classification_utilities.cfg.",
    "New file Config\\tc_config.txt holds ONLY [CREDENTIALS].",
    "Config\\tc_environment.cfg deleted.",
    "Build_All.bat reads TC_ROOT from the master cfg again (no tc file).",
    "Run_*.bat read TC_ROOT / TC_DATA / LOG_DIR from the master cfg and "
    "TC_USER / TC_PASS / TC_GROUP from tc_config.txt.",
    "The three exes resolve their credentials file as tc_config.txt "
    "(-tcconfig= argument or auto-discovery next to the -config= file / "
    "..\\..\\Config).",
])

# ------------------------------------------------------------- 3. architecture
h1("3. New Configuration Architecture")
add_code(
    "Config\\\\"
    "\n"
    r"  classification_utilities.cfg   DEVELOPER settings (ships with the exes,"
    "\n"
    r"                                 the user never edits it)"
    "\n"
    r"                                 [ENVIRONMENT] TC_ROOT, TC_DATA"
    "\n"
    r"                                 [BUILD] [LOGS] [EXTRACTION] [IMPORT]"
    "\n"
    r"                                 [DELETE] [RUN]"
    "\n"
    "\n"
    r"  tc_config.txt                  TEAMCENTER LOGIN ONLY (user-owned,"
    "\n"
    r"                                 the ONLY file the user edits)"
    "\n"
    r"                                 [CREDENTIALS] TC_USER, TC_PASS,"
    "\n"
    r"                                              TC_GROUP")
doc.add_paragraph("Who reads what:")
grid_table(
    ["Consumer", "classification_utilities.cfg", "tc_config.txt"],
    [
        ("Build_All.bat (developer)",
         "[BUILD] VCVARSALL/PLATFORM/CONFIGURATION, [ENVIRONMENT] TC_ROOT "
         "(generates TC_Root.props for msbuild)",
         "not read - building needs no login"),
        ("Run_Extraction.bat / Run_Import.bat / Run_Delete.bat (user)",
         "[ENVIRONMENT] TC_ROOT + TC_DATA (calls tc_profilevars.bat), "
         "[LOGS] LOG_DIR",
         "[CREDENTIALS] (passed to the exe via -u=/-p=/-g= and -tcconfig=)"),
        ("ClassificationExtraction.exe", "[EXTRACTION], [LOGS]", "[CREDENTIALS]"),
        ("ClassificationImport.exe / ClassificationDelete.exe",
         "[IMPORT] / [DELETE], [LOGS]", "[CREDENTIALS]"),
    ],
    (Inches(1.9), Inches(2.5), Inches(2.1)))

doc.add_paragraph()
note("Credential precedence in every exe: -u=/-p=/-g= command line overrides "
     "beat tc_config.txt. There is deliberately NO fallback to a "
     "[CREDENTIALS] section in the master cfg any more - the login lives in "
     "exactly one place. The ITK logic of all three utilities is untouched.")

# ------------------------------------------------------------- 4. file list
h1("4. Overview of All Changed / New Files")
grid_table(
    ["File", "Change", "Why"],
    [
        (r"Config\tc_config.txt", "NEW",
         "User-owned TC login: [CREDENTIALS] TC_USER / TC_PASS / TC_GROUP"),
        (r"Config\tc_environment.cfg", "DELETED",
         "Replaced by tc_config.txt; TC_ROOT / TC_DATA went back to the "
         "master cfg"),
        (r"Config\classification_utilities.cfg", "MODIFIED",
         "[ENVIRONMENT] TC_ROOT/TC_DATA restored; [CREDENTIALS] removed "
         "(pointer note added); header documents the split"),
        (r"Config\Build_All.bat", "MODIFIED",
         "TCFILE variable and its check removed; TC_ROOT read from the "
         "master cfg again (3-argument _cfg.bat call)"),
        (r"Config\Run_Extraction.bat, Run_Import.bat, Run_Delete.bat",
         "MODIFIED",
         "Credentials read from tc_config.txt; TC_ROOT/TC_DATA/LOG_DIR read "
         "from the master cfg; credential validation added; banner shows "
         "'TC credentials : ...tc_config.txt'"),
        (r"ClassificationExtraction\Source.cpp, "
         r"ClassificationImport\Source.cpp, ClassificationDelete\Source.cpp",
         "MODIFIED",
         "The Teamcenter credentials file is now named tc_config.txt "
         "(auto-discovery, error messages, usage text). Logic unchanged."),
        (r"ClassificationExtraction\ClassificationExtraction.vcxproj",
         "MODIFIED",
         "Comment block only: TC_ROOT now documented as coming from the "
         "master cfg"),
    ],
    (Inches(2.4), Inches(1.0), Inches(3.1)))

doc.add_page_break()

# ------------------------------------------------------------- change 1
h1("5. Change 1 - New file Config\\tc_config.txt")
doc.add_paragraph("This is the file the user edits after receiving the exes. "
                  "It is plain INI syntax, parsed by the same rules as the "
                  "master cfg. Full content:")
add_code(
    "[CREDENTIALS]\n"
    "; Your Teamcenter login\n"
    "TC_USER  = miguser\n"
    "TC_PASS  = miguser\n"
    "TC_GROUP = dba")
doc.add_paragraph("(The shipped file also carries a documentation header "
                  "explaining that this is the only file the user edits, and "
                  "the clear-text password warning - see the file itself.)")
expected(r"Config\tc_config.txt exists next to classification_utilities.cfg "
         r"and contains the [CREDENTIALS] section.")

# ------------------------------------------------------------- change 2
h1(r"6. Change 2 - classification_utilities.cfg: [ENVIRONMENT] restored, "
   "[CREDENTIALS] removed")
doc.add_paragraph(r"The [ENVIRONMENT] section is BACK in the master cfg "
                  "(developer-owned library settings):")
add_code(
    "[ENVIRONMENT]\n"
    "; Root of the Teamcenter installation (the folder that contains include\\,\n"
    "; include_cpp\\ and lib\\).  Example: D:\\Teamcenter_2312\\tc\n"
    "TC_ROOT  = D:\\Teamcenter_2312\\tc\n"
    "\n"
    "; TC data volume - the folder that contains tc_profilevars.bat.\n"
    "; Example: \\\\gptcdv-ep-atc04\\apps\\TCDATA\n"
    "TC_DATA  = \\\\gptcdv-ep-atc04\\apps\\TCDATA")
doc.add_paragraph("Where [CREDENTIALS] used to be, only a pointer note "
                  "remains:")
add_code(
    ";  NOTE: the Teamcenter login credentials ([CREDENTIALS] TC_USER/TC_PASS/\n"
    ";  TC_GROUP) live in Config\\tc_config.txt - the only file the user edits.")
doc.add_paragraph("The file header now states that this is the DEVELOPER "
                  "configuration file: the user edits ONLY tc_config.txt. "
                  "Everything else - [BUILD], [LOGS], [EXTRACTION], "
                  "[IMPORT], [DELETE], [RUN] - is unchanged.")
expected("grep -i \"TC_USER\\|TC_PASS\" classification_utilities.cfg returns "
         "no KEY = VALUE hits; TC_ROOT / TC_DATA are present again.")

# ------------------------------------------------------------- change 3
h1(r"7. Change 3 - Build_All.bat reads TC_ROOT from the master cfg again")
doc.add_paragraph("The tc_environment.cfg handling is removed:")
add_code(
    r"rem REMOVED:"
    "\n"
    r"set ""TCFILE=%CONFIGDIR%tc_environment.cfg"""
    "\n"
    "\n"
    r"rem REMOVED existence check:"
    "\n"
    r"if not exist ""%TCFILE%"" ("
    "\n"
    r"    echo ERROR : Teamcenter environment config file not found ..."
    "\n"
    r")"
    "\n"
    "\n"
    r"rem BEFORE (4th argument selected the tc file):"
    "\n"
    r"call ""%SCRIPTDIR%_cfg.bat"" ENVIRONMENT TC_ROOT TC_ROOT ""%TCFILE%"""
    "\n"
    "\n"
    r"rem AFTER (default master cfg):"
    "\n"
    r"call ""%SCRIPTDIR%_cfg.bat"" ENVIRONMENT TC_ROOT TC_ROOT")
doc.add_paragraph("The :err_cfg message points at the master cfg for both "
                  "keys:")
add_code(
    ":err_cfg\n"
    "echo ERROR : could not read [BUILD] VCVARSALL or [ENVIRONMENT] TC_ROOT\n"
    "echo         from %CFILE%\n"
    "goto :err_end")
doc.add_paragraph("Everything else (TC_Root.props generation, msbuild "
                  "invocation with /p:TC_ROOT=...) works as before.")
expected("Build_All.bat builds with only the master cfg present; "
         "tc_config.txt is NOT needed for building (a build needs no login).")

# ------------------------------------------------------------- change 4
h1("8. Change 4 - Run_Extraction.bat / Run_Import.bat / Run_Delete.bat")
doc.add_paragraph("All three wrappers change identically (shown for "
                  "Extraction; the other two differ only in exe name and "
                  "banner text):")
doc.add_heading("Step 8.1 - variables and existence checks", level=3)
add_code(
    r"set ""TCCFG=%SCRIPTDIR%tc_config.txt"""
    "\n"
    "\n"
    r"if not exist ""%TCCFG%"" ("
    "\n"
    r"    echo ERROR : Teamcenter credentials file not found : %TCCFG%"
    "\n"
    r"    echo         This is the ONLY file you need to edit - fill in your"
    "\n"
    r"    echo         Teamcenter login there, then run this bat again."
    "\n"
    r"    goto :end"
    "\n"
    r")")
doc.add_heading("Step 8.2 - split the reads: login vs library settings", level=3)
add_code(
    "rem ---- Teamcenter login : tc_config.txt (user-owned) ----------------\n"
    "call \"%SCRIPTDIR%_cfg.bat\" CREDENTIALS TC_USER  TC_USER  \"%TCCFG%\"\n"
    "call \"%SCRIPTDIR%_cfg.bat\" CREDENTIALS TC_PASS  TC_PASS  \"%TCCFG%\"\n"
    "call \"%SCRIPTDIR%_cfg.bat\" CREDENTIALS TC_GROUP TC_GROUP \"%TCCFG%\"\n"
    "\n"
    "rem ---- TC environment + function settings : master cfg ---------------\n"
    "call \"%SCRIPTDIR%_cfg.bat\" ENVIRONMENT TC_ROOT TC_ROOT\n"
    "call \"%SCRIPTDIR%_cfg.bat\" ENVIRONMENT TC_DATA TC_DATA\n"
    "call \"%SCRIPTDIR%_cfg.bat\" LOGS LOG_DIR LOG_DIR")
doc.add_heading("Step 8.3 - validation with the right owner in the message",
                level=3)
add_code(
    "rem missing login -> the USER fixes it (tc_config.txt)\n"
    "if \"%TC_USER%\"==\"\" (\n"
    "    echo ERROR : [CREDENTIALS] TC_USER missing in %TCCFG%\n"
    "    goto :end\n"
    ")\n"
    "rem ... same for TC_PASS and TC_GROUP ...\n"
    "\n"
    "rem missing library settings -> the DEVELOPER fixes it (master cfg)\n"
    "if \"%TC_ROOT%\"==\"\" (\n"
    "    echo ERROR : [ENVIRONMENT] TC_ROOT missing in %CFILE% - contact the developer.\n"
    "    goto :end\n"
    ")\n"
    "if not exist \"%TC_DATA%\\tc_profilevars.bat\" (\n"
    "    echo ERROR : %TC_DATA%\\tc_profilevars.bat not found.\n"
    "    echo         Check [ENVIRONMENT] TC_DATA in %CFILE% - contact the developer.\n"
    "    goto :end\n"
    ")\n"
    "call \"%TC_DATA%\\tc_profilevars.bat\"")
doc.add_heading("Step 8.4 - banner + invocation", level=3)
add_code(
    "echo    TC credentials : %TCCFG%\n"
    "\n"
    "rem the exe gets the login twice: as overrides AND via -tcconfig=,\n"
    "rem so it also works when the user runs the exe directly\n"
    "\"%EXE%\" -u=\"%TC_USER%\" -p=\"%TC_PASS%\" -g=\"%TC_GROUP%\" -config=\"%CFILE%\" -tcconfig=\"%TCCFG%\"")
expected("Each Run_*.bat prints 'TC credentials : ...tc_config.txt' in its "
         "banner, fails fast with a clear message when the login file or "
         "tc_profilevars.bat is missing, and never asks the user to touch "
         "classification_utilities.cfg.")

# ------------------------------------------------------------- change 5
h1("9. Change 5 - the three exes read credentials from tc_config.txt")
doc.add_paragraph("All three utilities keep the argument "
                  "-tcconfig=<path to tc_config.txt>. When the argument is "
                  "not given, each exe auto-discovers the file: first next "
                  "to the -config= file, then in ..\\..\\Config relative to "
                  "the exe, then .\\Config, then the current directory. "
                  "Precedence: command line -u=/-p=/-g= beats tc_config.txt. "
                  "In this rework only the FILE NAME and the messages "
                  "changed - the resolution logic existed already.")
bullets([
    r"ClassificationExtraction\Source.cpp: findTcConfigFile() searches for "
    r""""tc_config.txt" (was "tc_environment.cfg"); the hard error (exit 2) "
    "when the file is missing stays, and its message now says: Create "
    "Config\\tc_config.txt and fill in [CREDENTIALS].""",
    r"ClassificationImport\Source.cpp + ClassificationDelete\Source.cpp: the "
    "default path derived from -config= becomes "
    "<config dir>\\tc_config.txt; the missing-credentials error points at "
    "tc_config.txt; both keep their tolerant load (missing file only fails "
    "at the credentials check).",
    "All three displayUsage() texts now show [-tcconfig=<tc_config.txt "
    "path>] and the sentence 'login comes from the [CREDENTIALS] section of "
    "tc_config.txt'.",
])
add_code(
    "/* Extraction - findTcConfigFile */\n"
    "const string sConfigName = \"tc_config.txt\";\n"
    "\n"
    "/* Import / Delete - default next to the -config= file */\n"
    "sTcConfigPath = sConfigFile.substr(0, nLastSlash + 1) + \"tc_config.txt\";\n"
    "\n"
    "/* credentials in all three exes (unchanged logic) */\n"
    "if (sUserId.empty()) sUserId = oTcCfg.getString(\"CREDENTIALS\", \"TC_USER\", \"\");\n"
    "if (sPwd.empty())    sPwd    = oTcCfg.getString(\"CREDENTIALS\", \"TC_PASS\", \"\");\n"
    "if (sGroup.empty())  sGroup  = oTcCfg.getString(\"CREDENTIALS\", \"TC_GROUP\", \"\");")
expected("Each exe logs in with the credentials from tc_config.txt; "
         "'Using Teamcenter environment file : ...tc_config.txt' is printed "
         "when the file was found.")

doc.add_page_break()

# ------------------------------------------------------------- 10. delivery
h1("10. Delivery Package - What the Developer Ships")
grid_table(
    ["Delivered item", "Purpose"],
    [
        (r"x64\Release\ClassificationExtraction.exe", "Extraction utility"),
        (r"x64\Release\ClassificationImport.exe", "Import utility"),
        (r"x64\Release\ClassificationDelete.exe", "Delete/Validation utility"),
        (r"Config\classification_utilities.cfg",
         "Developer settings (TC library paths + function settings) - works "
         "as shipped, user never edits it"),
        (r"Config\tc_config.txt",
         "Template - the user fills in their TC login ONLY"),
        (r"Config\Run.bat, Run_*.bat, _cfg.bat", "Runtime wrappers"),
    ],
    (Inches(2.9), Inches(3.6)))

warn("Do NOT ship Config\\TC_Root.props (a build-machine artifact) and reset "
     "tc_config.txt to placeholder/example credentials before packaging. "
     "Do NOT ship the developer's password. The exes contain no Teamcenter "
     "path dependency - paths come from the master cfg at runtime.")

doc.add_heading("User workflow after delivery (this is the whole story)",
                level=2)
numbered([
    "Copy the package to any folder on the target machine.",
    "Open Config\\tc_config.txt.",
    "Set [CREDENTIALS] TC_USER / TC_PASS / TC_GROUP to the local Teamcenter "
    "login.",
    "Save the file. (Nothing else - do not edit "
    "classification_utilities.cfg.)",
    "Double-click Config\\Run.bat (or the Run_*.bat of the utility).",
])
expected("Run.bat resolves [RUN] FUNCTION, the wrapper shows the "
         "'TC credentials : ...tc_config.txt' banner, calls "
         "tc_profilevars.bat from [ENVIRONMENT] TC_DATA, starts the exe with "
         "the login from tc_config.txt and the utility performs its task - "
         "with exactly ONE file edited by the user.")

# ------------------------------------------------------------- 11. build
h1("11. Build Procedure (developer machine)")
numbered([
    "Ensure [ENVIRONMENT] TC_ROOT and [BUILD] VCVARSALL in "
    "classification_utilities.cfg point at this machine.",
    "Run Config\\Build_All.bat (optionally ALL clean) - or "
    "Build_All.bat EXTRACTION / IMPORT / DELETE for a single utility.",
    "Build_All.bat reads TC_ROOT from the master cfg, regenerates "
    "TC_Root.props, calls vcvarsall + msbuild.",
    "Output lands in x64\\Release\\*.exe as before.",
])
doc.add_paragraph(
    "Visual Studio IDE builds also work: after one Build_All.bat run, "
    "TC_Root.props exists and the projects resolve $(TC_ROOT); otherwise set "
    "the TC_ROOT environment variable before opening VS.")

# ------------------------------------------------------------- 12. checklist
h1("12. Verification Checklist")
grid_table(
    ["Check", "Expected"],
    [
        ("classification_utilities.cfg",
         "has [ENVIRONMENT] TC_ROOT/TC_DATA; has NO TC_USER/TC_PASS/TC_GROUP"),
        ("tc_config.txt exists with [CREDENTIALS]",
         "file present next to the master cfg"),
        ("call _cfg.bat CREDENTIALS TC_USER R tc_config.txt",
         "R = the TC_USER value"),
        ("call _cfg.bat ENVIRONMENT TC_ROOT R (3-argument form)",
         "R = the TC_ROOT value from the master cfg"),
        ("Build_All.bat on a TC machine",
         "BUILD OK, exes in x64\\Release"),
        ("Rename tc_config.txt away, run Run_Extraction.bat",
         "'Teamcenter credentials file not found' error - no half-run"),
        ("Empty TC_PASS in tc_config.txt, run again",
         "'[CREDENTIALS] TC_PASS missing in ...tc_config.txt' error"),
        ("Restore the file, run Run_Extraction.bat",
         "'TC credentials :' banner, tc_profilevars.bat called, exe reports "
         "'Using Teamcenter environment file : ...tc_config.txt', login OK"),
        ("Extraction -h / Import -h / Delete -h",
         "Usage shows [-tcconfig=<tc_config.txt path>]"),
        ("Wrong TC_DATA (folder without tc_profilevars.bat)",
         "Clear error mentioning tc_profilevars.bat and the master cfg"),
    ],
    (Inches(3.0), Inches(3.5)))

# ------------------------------------------------------------- 13. trouble
h1("13. Troubleshooting")
grid_table(
    ["Symptom", "Cause / Fix"],
    [
        ("'Teamcenter credentials file not found'",
         "tc_config.txt is not next to classification_utilities.cfg and no "
         "-tcconfig= was passed. Place the file or pass the path."),
        ("'[CREDENTIALS] TC_xxx missing in ...tc_config.txt'",
         "The user has not filled in the login - edit tc_config.txt."),
        ("Login failed (exit 3/4)",
         "Wrong credentials in tc_config.txt, or the TC_DATA in the master "
         "cfg points to another TC environment than the user/password "
         "belongs to (developer fixes TC_DATA)."),
        ("Build: 'TC_ROOT does not look like a Teamcenter installation'",
         "[ENVIRONMENT] TC_ROOT in classification_utilities.cfg wrong or "
         "empty - point it at the folder containing "
         "include\\tc\\tc_startup.h."),
        ("Run: 'tc_profilevars.bat not found'",
         "[ENVIRONMENT] TC_DATA wrong in the master cfg - developer fix."),
        ("'cannot open include file tc/tc_startup.h' in VS",
         "TC_Root.props missing/stale and no TC_ROOT env var - run "
         "Build_All.bat once, then reopen the solution."),
        ("Credentials ignored",
         "Precedence is -u=/-p=/-g= > tc_config.txt. Check for stray -u= "
         "flags in custom callers."),
        ("User edited classification_utilities.cfg and broke the run",
         "Restore the file from the delivery package - users never need to "
         "change it."),
    ],
    (Inches(2.7), Inches(3.8)))

# ------------------------------------------------------------- 14. rollback
h1("14. Rollback")
bullets([
    "Git: revert the files listed in section 4 (git checkout -- <files>) "
    "and delete Config\\tc_config.txt.",
    "Manual rollback to the WRONG split (not recommended): recreate "
    "tc_environment.cfg with [ENVIRONMENT] + [CREDENTIALS], revert "
    "Build_All.bat + Run_*.bat + the three Source.cpp files to the "
    "tc_environment.cfg naming.",
    "Manual rollback to the ORIGINAL single-file layout: put [CREDENTIALS] "
    "back into classification_utilities.cfg, make the Run_*.bat read the "
    "login from the master cfg, remove -tcconfig= from the invocation lines "
    "and the three Source.cpp files.",
    "Temporary coexistence: keeping the -tcconfig= argument while also "
    "shipping a [CREDENTIALS] section in the master cfg does NOT work with "
    "this code - the exes no longer read credentials from the master cfg.",
])

# ------------------------------------------------------------- save
try:
    doc.save(OUT_PATH)
    print("Written:", OUT_PATH)
except PermissionError:
    alt = OUT_PATH.replace(".docx", " (new).docx")
    doc.save(alt)
    print("Original docx is open in Word - written instead:", alt)
